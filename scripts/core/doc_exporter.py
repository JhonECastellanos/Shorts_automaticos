"""Exportador de documentos TXT y DOCX por short."""

from pathlib import Path

from .utils import (
    ProgressCallback, noop_progress, get_episode_dir, load_json,
    srt_time_to_seconds, seconds_to_srt_time,
)


def export_documents(
    *,
    project: str,
    episode: str,
    output_base: Path | None = None,
    root: Path,
    on_progress: ProgressCallback = noop_progress,
) -> list[dict]:
    ep_dir = get_episode_dir(root, project, episode, create=False)
    analysis_dir = ep_dir / "analysis"

    moments_path = analysis_dir / "moments.json"
    ranked_path = analysis_dir / "moments_ranked.json"

    if ranked_path.exists():
        moments = load_json(ranked_path)
    elif moments_path.exists():
        moments = load_json(moments_path)
    else:
        raise FileNotFoundError("No se encontró moments.json. Ejecuta analyze primero.")

    # Cargar SRT completo para extraer texto
    srt_files = list((ep_dir / "transcripts").glob("*.srt"))
    full_srt_blocks = _parse_srt(srt_files[0]) if srt_files else []

    # Directorio de salida
    if output_base is None:
        output_base = ep_dir / "output"

    results = []
    for i, moment in enumerate(moments):
        short_num = moment.get("rank", i + 1)
        short_dir = output_base / f"short_{short_num:02d}"
        short_dir.mkdir(parents=True, exist_ok=True)

        pct = (i / len(moments)) * 90
        on_progress("doc_export", f"Exportando docs short_{short_num:02d}", pct)

        # Extraer texto del rango
        text_lines = _extract_text_for_range(full_srt_blocks, moment["start"], moment["end"])

        # TXT
        txt_path = short_dir / f"short_{short_num:02d}.txt"
        txt_content = _build_txt(moment, text_lines, short_num)
        txt_path.write_text(txt_content, encoding="utf-8")

        # DOCX
        docx_path = short_dir / f"short_{short_num:02d}.docx"
        _build_docx(moment, text_lines, short_num, docx_path)

        results.append({"short_num": short_num, "txt": str(txt_path), "docx": str(docx_path)})

    on_progress("doc_export", f"{len(results)} documentos exportados", 100)
    return results


def _parse_srt(srt_path: Path) -> list[dict]:
    content = srt_path.read_text(encoding="utf-8")
    blocks = content.strip().split("\n\n")
    parsed = []
    for block in blocks:
        lines = block.strip().split("\n")
        if len(lines) < 3:
            continue
        time_line = lines[1]
        parts = time_line.split(" --> ")
        if len(parts) != 2:
            continue
        parsed.append({
            "start": srt_time_to_seconds(parts[0]),
            "end": srt_time_to_seconds(parts[1]),
            "text": " ".join(lines[2:]),
        })
    return parsed


def _extract_text_for_range(blocks: list[dict], start: float, end: float) -> list[str]:
    lines = []
    for b in blocks:
        if b["end"] <= start or b["start"] >= end:
            continue
        lines.append(b["text"])
    return lines


def _build_txt(moment: dict, text_lines: list[str], short_num: int) -> str:
    duration = moment.get("end", 0) - moment.get("start", 0)
    lines = [
        f"SHORT #{short_num}",
        f"{'=' * 40}",
        f"Tema: {moment.get('topic', 'N/A')}",
        f"Hook: {moment.get('hook', 'N/A')}",
        f"Score: {moment.get('score', 'N/A')}",
        f"Trend Score: {moment.get('trend_score', 'N/A')}",
        f"Duración: {duration:.1f}s ({seconds_to_srt_time(moment.get('start', 0))} → {seconds_to_srt_time(moment.get('end', 0))})",
        f"Speaker: {moment.get('dominant_speaker', 'N/A')}",
        f"Razón: {moment.get('reason', 'N/A')}",
        "",
        "TRANSCRIPCIÓN:",
        "-" * 40,
    ]
    lines.extend(text_lines)
    return "\n".join(lines)


def _build_docx(moment: dict, text_lines: list[str], short_num: int, out_path: Path) -> None:
    try:
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    except ImportError:
        # Si python-docx no está instalado, solo escribe TXT
        return

    doc = Document()

    # Título
    title = doc.add_heading(f"Short #{short_num}", level=1)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # Metadata tabla
    duration = moment.get("end", 0) - moment.get("start", 0)
    table = doc.add_table(rows=6, cols=2, style="Light List Accent 1")
    fields = [
        ("Tema", moment.get("topic", "N/A")),
        ("Hook", moment.get("hook", "N/A")),
        ("Score IA", str(moment.get("score", "N/A"))),
        ("Trend Score", str(moment.get("trend_score", "N/A"))),
        ("Duración", f"{duration:.1f}s"),
        ("Speaker", moment.get("dominant_speaker", "N/A")),
    ]
    for i, (label, value) in enumerate(fields):
        table.cell(i, 0).text = label
        table.cell(i, 1).text = value

    # Razón
    doc.add_heading("Por qué funciona", level=2)
    doc.add_paragraph(moment.get("reason", "N/A"))

    # Transcripción
    doc.add_heading("Transcripción", level=2)
    for line in text_lines:
        doc.add_paragraph(line)

    doc.save(str(out_path))
